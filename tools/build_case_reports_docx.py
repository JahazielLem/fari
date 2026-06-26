#!/usr/bin/env python3
"""Build separate Word editions of the canonical FARI case reports."""

from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt

from build_reference_docx import (
    BLUE,
    INK,
    MUTED,
    add_callout,
    add_result_banner,
    configure_running_furniture,
    configure_styles,
    markdown_to_doc,
    set_run_font,
)


ROOT = Path(__file__).resolve().parents[1]
SPEC_REPORTS = ROOT / "spec" / "current" / "reports"
OUTPUT_DIR = ROOT / "artifacts" / "reports"

REPORTS = [
    (
        SPEC_REPORTS / "01-library-fuzzing-report.md",
        OUTPUT_DIR / "FARI-Case-01-Library-Fuzzing-Report.docx",
        "Investigation Report",
        "OSDLP Telecommand Library Fuzzing",
        "Auditor evidence: cases/01_Lib.md",
        "7 June 2026",
    ),
    (
        SPEC_REPORTS / "02-usb-interface-report.md",
        OUTPUT_DIR / "FARI-Case-02-USB-Interface-Report.docx",
        "Investigation Report",
        "USB-Simulated Telecommand Interface",
        "Auditor evidence: cases/02_USB.md",
        "7 June 2026",
    ),
    (
        SPEC_REPORTS / "03-usb-apid-underflow-report.md",
        OUTPUT_DIR / "FARI-Case-03-USB-APID-Underflow-Report.docx",
        "Investigation Report",
        "USB APID 0x06 Undersized Telecommand",
        "Auditor evidence: cases/03_USB.md and usb_tc_apid_underflow_cmd.png",
        "8 June 2026",
    ),
    (
        SPEC_REPORTS / "04-spacecanbus-spoofing-replay-report.md",
        OUTPUT_DIR / "FARI-Case-04-SpaceCAN-Spoofing-Replay-Report.docx",
        "Investigation Report",
        "SpaceCAN Spoofing and Replay",
        "Auditor evidence: cases/04_Spacecanbus.md and supplied screenshots",
        "8 June 2026",
    ),
    (
        SPEC_REPORTS / "FARI-CONSOLIDATED-ASSESSMENT-REPORT.md",
        OUTPUT_DIR / "FARI-Consolidated-Assessment-Report.docx",
        "Consolidated Assessment Report",
        "Pilot Validation Portfolio",
        "Source material: cases/01_Lib.md through cases/04_Spacecanbus.md",
        "8 June 2026",
    ),
]


def add_report_cover(doc, report_label, title, evidence_source, report_date, conclusion):
    for _ in range(5):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("FARI")
    set_run_font(run, size=30, color=INK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(report_label)
    set_run_font(run, size=18, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    run = p.add_run(title)
    set_run_font(run, size=15, color=MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f"Report maturity: Provisional  |  {report_date}")
    set_run_font(run, size=10.5, color=MUTED, bold=True)

    for _ in range(2):
        doc.add_paragraph()

    add_result_banner(doc, conclusion, "FARI qualitative executive conclusion")

    add_callout(
        doc,
        "External auditors do not need FARI knowledge or forms. Supplied source "
        "material remains in cases/; the Report Author creates all FARI "
        "analysis, conclusions, and follow-up actions.",
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    run = p.add_run(evidence_source)
    set_run_font(run, size=9.5, color=MUTED, italic=True)
    p.add_run().add_break(WD_BREAK.PAGE)


def trim_trailing_empty_paragraphs(doc):
    while doc.paragraphs and not doc.paragraphs[-1].text.strip():
        paragraph = doc.paragraphs[-1]
        paragraph._element.getparent().remove(paragraph._element)


def build_report(source, output, report_label, title, evidence_source, report_date):
    doc = Document()
    configure_styles(doc)
    configure_running_furniture(doc.sections[0])
    source_text = source.read_text(encoding="utf-8")
    match = re.search(
        r"\|\s*(?:Overall )?Conclusion\s*\|\s*\*\*([^*]+)\*\*\s*\|",
        source_text,
        re.IGNORECASE,
    )
    conclusion = match.group(1) if match else "Not Assessed"
    add_report_cover(
        doc,
        report_label,
        title,
        evidence_source,
        report_date,
        conclusion,
    )
    markdown_to_doc(doc, source, page_break=False)
    trim_trailing_empty_paragraphs(doc)
    doc.core_properties.title = f"FARI {report_label}: {title}"
    doc.core_properties.subject = f"FARI {report_label}"
    doc.core_properties.author = "FARI Report Author"
    doc.core_properties.keywords = "FARI, space cybersecurity, investigation, SPARTA"
    doc.save(output)
    print(output)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for args in REPORTS:
        build_report(*args)


if __name__ == "__main__":
    main()
