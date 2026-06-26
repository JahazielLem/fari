#!/usr/bin/env python3
"""Build the FARI Word reference guide from the canonical Markdown documents."""

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SPEC_CURRENT_ROOT = ROOT / "spec" / "current"
OUTPUT = ROOT / "artifacts" / "FARI-Reference-Guide.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
RESULT_STYLES = {
    "meets": ("MEETS", "15803D", WHITE),
    "does not meet": ("DOES NOT MEET", "B42318", WHITE),
    "inconclusive": ("INCONCLUSIVE", "B54708", WHITE),
    "not assessed": ("NOT ASSESSED", "475467", WHITE),
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def new_numbering_instance(doc):
    numbering = doc.part.numbering_part.element
    existing = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
        if node.get(qn("w:numId"))
    ]
    num_id = max(existing, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), "0")
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return num_id


def set_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)


def set_table_geometry(table, widths):
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline_runs(paragraph, text, size=11, color=None):
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Courier New", size=max(9, size - 1), color=DARK_BLUE)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=size, color=color, italic=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def configure_running_furniture(section):
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(0)
    run = header_p.add_run("FARI  |  Framework for Aerospace Research and Investigation")
    set_run_font(run, size=9, color=MUTED, bold=True)
    add_page_number(section.footer.paragraphs[0])


def add_cover(doc):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("FARI")
    set_run_font(run, size=32, color=INK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("Framework for Aerospace Research and Investigation")
    set_run_font(run, size=18, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    run = p.add_run("Methodology, Integration Profile, Examples, and Process Validation")
    set_run_font(run, size=13, color=MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("Version 0.3 Draft  |  8 June 2026")
    set_run_font(run, size=10.5, color=MUTED, bold=True)

    for _ in range(3):
        doc.add_paragraph()

    table = doc.add_table(rows=1, cols=4)
    set_table_geometry(table, [2340, 2340, 2340, 2340])
    set_repeat_table_header(table.rows[0])
    labels = [
        ("FRAME", "Mission, scope, authorization"),
        ("ACQUIRE", "Evidence intake and sufficiency"),
        ("RELATE", "SPARTA, findings, conclusions"),
        ("INFORM", "Result cards, consolidation, action"),
    ]
    for idx, (label, detail) in enumerate(labels):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_BLUE if idx % 2 == 0 else LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(label)
        set_run_font(r, size=11, color=INK, bold=True)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(detail)
        set_run_font(r, size=9, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "Canonical source: repository Markdown and JSON schema. "
        "This Word guide is a derived reference edition."
    )
    set_run_font(run, size=9.5, color=MUTED, italic=True)
    p.add_run().add_break(WD_BREAK.PAGE)


def add_callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline_runs(p, text, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def result_style(text):
    normalized = re.sub(r"[*_`\[\]]", "", text).strip().lower()
    for key, style in RESULT_STYLES.items():
        if key in normalized:
            return style
    return RESULT_STYLES["not assessed"]


def add_result_banner(doc, conclusion, detail=None):
    label, fill, text_color = result_style(conclusion)
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0 if not detail else 4)
    run = p.add_run(label)
    set_run_font(run, size=18, color=text_color, bold=True)
    if detail:
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(detail)
        set_run_font(run, size=9.5, color=text_color, bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_markdown_table(doc, rows):
    max_cols = max(len(row) for row in rows)
    clean_rows = []
    for row in rows:
        if all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in row):
            continue
        clean_rows.append(row + [""] * (max_cols - len(row)))
    if not clean_rows:
        return

    table = doc.add_table(rows=len(clean_rows), cols=max_cols)
    widths = [9360 // max_cols] * max_cols
    widths[-1] += 9360 - sum(widths)
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(clean_rows):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            if r_idx == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            add_inline_runs(p, text.strip(), size=9.2, color=INK)
            if r_idx == 0:
                for run in p.runs:
                    run.bold = True
        label = clean_rows[r_idx][0].strip().lower()
        if label in {"conclusion", "overall conclusion"} and len(clean_rows[r_idx]) > 1:
            _status_label, fill, text_color = result_style(clean_rows[r_idx][1])
            for cell in table.rows[r_idx].cells:
                set_cell_shading(cell, fill)
                set_cell_margins(cell, top=140, start=150, bottom=140, end=150)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_run_font(
                            run,
                            size=11.5,
                            color=text_color,
                            bold=True,
                        )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)


def parse_table_line(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def markdown_to_doc(doc, path, title_override=None, page_break=True):
    lines = path.read_text(encoding="utf-8").splitlines()
    if page_break:
        doc.add_page_break()
    if title_override:
        p = doc.add_paragraph(style="Heading 1")
        add_inline_runs(p, title_override, size=16, color=BLUE)

    idx = 0
    paragraph_lines = []
    in_code = False
    code_language = ""
    code_lines = []
    active_num_id = None

    def flush_paragraph():
        nonlocal paragraph_lines
        if paragraph_lines:
            p = doc.add_paragraph()
            add_inline_runs(p, " ".join(x.strip() for x in paragraph_lines))
            paragraph_lines = []

    while idx < len(lines):
        line = lines[idx]

        if line.startswith("```"):
            flush_paragraph()
            if not in_code:
                in_code = True
                code_language = line[3:].strip()
                code_lines = []
            else:
                if code_language == "mermaid":
                    add_callout(
                        doc,
                        "Diagram source is available in the canonical Markdown edition.",
                    )
                else:
                    table = doc.add_table(rows=1, cols=1)
                    set_table_geometry(table, [9360])
                    set_repeat_table_header(table.rows[0])
                    cell = table.cell(0, 0)
                    set_cell_shading(cell, LIGHT_GRAY)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(0)
                    r = p.add_run("\n".join(code_lines))
                    set_run_font(r, name="Courier New", size=8.5, color=INK)
                in_code = False
                code_language = ""
                code_lines = []
            idx += 1
            continue

        if in_code:
            code_lines.append(line)
            idx += 1
            continue

        if line.startswith("|") and idx + 1 < len(lines) and lines[idx + 1].startswith("|"):
            flush_paragraph()
            table_rows = []
            while idx < len(lines) and lines[idx].startswith("|"):
                table_rows.append(parse_table_line(lines[idx]))
                idx += 1
            add_markdown_table(doc, table_rows)
            continue

        if line.strip() == "<!-- pagebreak -->":
            flush_paragraph()
            doc.add_page_break()
            idx += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            flush_paragraph()
            level = min(len(heading.group(1)), 3)
            text = heading.group(2)
            if level == 1 and (text.startswith("Framework for") or text.startswith("FARI ")):
                idx += 1
                continue
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline_runs(p, text, size={1: 16, 2: 13, 3: 12}[level], color=BLUE if level < 3 else DARK_BLUE)
            active_num_id = None
            idx += 1
            continue

        bullet = re.match(r"^\s*-\s+(.+)$", line)
        numbered = re.match(r"^\s*\d+\.\s+(.+)$", line)
        if bullet or numbered:
            flush_paragraph()
            item_text = (bullet or numbered).group(1)
            next_idx = idx + 1
            while (
                next_idx < len(lines)
                and re.match(r"^\s{2,}\S", lines[next_idx])
                and not re.match(r"^\s*[-*]\s+", lines[next_idx])
                and not re.match(r"^\s*\d+\.\s+", lines[next_idx])
            ):
                item_text += " " + lines[next_idx].strip()
                next_idx += 1
            if numbered:
                if active_num_id is None:
                    active_num_id = new_numbering_instance(doc)
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.25
                set_numbering(p, active_num_id)
            else:
                active_num_id = None
                p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, item_text)
            idx = next_idx
            continue

        if line.startswith(">"):
            flush_paragraph()
            quote_lines = []
            while idx < len(lines) and lines[idx].startswith(">"):
                quote_lines.append(lines[idx].lstrip("> "))
                idx += 1
            add_callout(doc, " ".join(quote_lines))
            continue

        if line.strip() == "---":
            flush_paragraph()
            idx += 1
            continue

        if not line.strip():
            flush_paragraph()
        else:
            active_num_id = None
            paragraph_lines.append(line)
        idx += 1

    flush_paragraph()


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    configure_running_furniture(doc.sections[0])
    add_cover(doc)

    add_callout(
        doc,
        "FARI is a top-level integration and assurance framework. It does not "
        "replace technical audit methods or burden external auditors; it makes "
        "their source material, context, coverage, and qualitative executive "
        "conclusions interoperable.",
    )

    for path, title in (
        (SPEC_CURRENT_ROOT / "specification" / "FARI-SPECIFICATION.md", None),
        (SPEC_CURRENT_ROOT / "specification" / "FARI-SPD5-COMPANION.md", "Appendix A: SPD-5 Companion Profile"),
        (ROOT / "examples" / "01-minimal-source-review.md", "Appendix C: Progressive Examples"),
        (ROOT / "examples" / "02-qemu-firmware-assessment.md", None),
        (ROOT / "examples" / "03-mission-wide-assessment.md", None),
        (SPEC_CURRENT_ROOT / "reports" / "FARI-PROCESS-VALIDATION.md", "Appendix D: Process Validation"),
    ):
        markdown_to_doc(doc, path, title_override=title, page_break=path.name != "FARI-METHODOLOGY.md")

    doc.core_properties.title = "FARI Framework for Aerospace Research and Investigation"
    doc.core_properties.subject = "Methodology, SPARTA integration, examples, process validation, and web requirements"
    doc.core_properties.author = "FARI Project"
    doc.core_properties.keywords = "FARI, space cybersecurity, SPARTA, assurance, qualitative conclusion"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
