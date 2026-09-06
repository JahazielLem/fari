#!/usr/bin/env python3
"""Build the versioned FARI documentation release as DOCX intermediates."""

from __future__ import annotations

import re
import shutil
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt

from tools.build_reference_docx import (
    BLUE,
    INK,
    MUTED,
    add_callout,
    add_result_banner,
    configure_running_furniture,
    configure_styles,
    markdown_to_doc,
    set_run_font,
)
from webapp.meta import FARI_GENERATED_WITH, FARI_VERSION


ROOT = Path(__file__).resolve().parents[1]
VERSION_TAG = f"v{FARI_VERSION}"
SPEC_ROOT = ROOT / "spec"
SPEC_CURRENT_ROOT = SPEC_ROOT / "current"
RELEASE_ROOT = SPEC_ROOT / "releases" / VERSION_TAG
SOURCE_ROOT = RELEASE_ROOT / "_sources"
BUILD_ROOT = RELEASE_ROOT / "_build" / "docx"
QA_ROOT = RELEASE_ROOT / "_qa"

SPEC_DIR = RELEASE_ROOT / "specification"
REPORTS_DIR = RELEASE_ROOT / "reports"
TEMPLATE_DIR = RELEASE_ROOT / "manual-template"


def _existing_release_sources() -> list[Path]:
    roots = []
    releases_root = SPEC_ROOT / "releases"
    if not releases_root.exists():
        return roots
    for path in sorted(releases_root.glob("v*/_sources"), reverse=True):
        if path.resolve() != SOURCE_ROOT.resolve():
            roots.append(path)
    return roots


def _candidate_paths(relative: str, *extra: Path) -> list[Path]:
    candidates = [path for path in extra if path]
    for root in _existing_release_sources():
        candidates.append(root / relative)
    candidates.append(SOURCE_ROOT / relative)
    return candidates


CURRENT_MARKDOWN_SOURCES = {
    "specification/FARI-SPECIFICATION.md": _candidate_paths(
        "specification/FARI-SPECIFICATION.md",
        SPEC_CURRENT_ROOT / "specification" / "FARI-SPECIFICATION.md",
        ROOT / "docs" / "FARI-SPECIFICATION.md",
    ),
    "reports/FARI-SCENARIO-TEST-CATALOG.md": _candidate_paths(
        "reports/FARI-SCENARIO-TEST-CATALOG.md",
        SPEC_CURRENT_ROOT / "reports" / "FARI-SCENARIO-TEST-CATALOG.md",
        ROOT / "docs" / "FARI-SCENARIO-TEST-CATALOG.md",
    ),
    "reports/01-library-fuzzing-report.md": _candidate_paths(
        "reports/01-library-fuzzing-report.md",
        SPEC_CURRENT_ROOT / "reports" / "01-library-fuzzing-report.md",
        ROOT / "reports" / "01-library-fuzzing-report.md",
    ),
    "reports/02-usb-interface-report.md": _candidate_paths(
        "reports/02-usb-interface-report.md",
        SPEC_CURRENT_ROOT / "reports" / "02-usb-interface-report.md",
        ROOT / "reports" / "02-usb-interface-report.md",
    ),
    "reports/03-usb-apid-underflow-report.md": _candidate_paths(
        "reports/03-usb-apid-underflow-report.md",
        SPEC_CURRENT_ROOT / "reports" / "03-usb-apid-underflow-report.md",
        ROOT / "reports" / "03-usb-apid-underflow-report.md",
    ),
    "reports/04-spacecanbus-spoofing-replay-report.md": _candidate_paths(
        "reports/04-spacecanbus-spoofing-replay-report.md",
        SPEC_CURRENT_ROOT / "reports" / "04-spacecanbus-spoofing-replay-report.md",
        ROOT / "reports" / "04-spacecanbus-spoofing-replay-report.md",
    ),
    "reports/FARI-CONSOLIDATED-ASSESSMENT-REPORT.md": _candidate_paths(
        "reports/FARI-CONSOLIDATED-ASSESSMENT-REPORT.md",
        SPEC_CURRENT_ROOT / "reports" / "FARI-CONSOLIDATED-ASSESSMENT-REPORT.md",
        ROOT / "reports" / "FARI-CONSOLIDATED-ASSESSMENT-REPORT.md",
    ),
    "reports/FARI-PROCESS-VALIDATION.md": _candidate_paths(
        "reports/FARI-PROCESS-VALIDATION.md",
        SPEC_CURRENT_ROOT / "reports" / "FARI-PROCESS-VALIDATION.md",
        ROOT / "reports" / "FARI-PROCESS-VALIDATION.md",
    ),
    "templates/FARI-Manual-Template.md": _candidate_paths(
        "templates/FARI-Manual-Template.md",
        SPEC_CURRENT_ROOT / "templates" / "FARI-Manual-Template.md",
        ROOT / "templates" / "FARI-Manual-Template.md",
    ),
}

LEGACY_DOCX_SOURCES = [
    ROOT / "artifacts" / "FARI-Reference-Guide.docx",
    ROOT / "artifacts" / "guides" / "FARI-Datasheet.docx",
    ROOT / "artifacts" / "guides" / "FARI-Manual-Walkthrough-SpaceCAN.docx",
    ROOT / "artifacts" / "reports" / "FARI-Case-01-Library-Fuzzing-Report.docx",
    ROOT / "artifacts" / "reports" / "FARI-Case-02-USB-Interface-Report.docx",
    ROOT / "artifacts" / "reports" / "FARI-Case-03-USB-APID-Underflow-Report.docx",
    ROOT / "artifacts" / "reports" / "FARI-Case-04-SpaceCAN-Spoofing-Replay-Report.docx",
    ROOT / "artifacts" / "reports" / "FARI-Consolidated-Assessment-Report.docx",
    ROOT / "artifacts" / "templates" / "FARI-Optional-Technical-Finding-Submission-Form.docx",
]

REPORT_BUILD_PLAN = [
    (
        SOURCE_ROOT / "reports" / "01-library-fuzzing-report.md",
        REPORTS_DIR / f"FARI-Case-01-Library-Fuzzing-Report-{VERSION_TAG}.pdf",
        "Example Investigation Report",
        "OSDLP Telecommand Library Fuzzing",
    ),
    (
        SOURCE_ROOT / "reports" / "02-usb-interface-report.md",
        REPORTS_DIR / f"FARI-Case-02-USB-Interface-Report-{VERSION_TAG}.pdf",
        "Example Investigation Report",
        "USB-Simulated Telecommand Interface",
    ),
    (
        SOURCE_ROOT / "reports" / "03-usb-apid-underflow-report.md",
        REPORTS_DIR / f"FARI-Case-03-USB-APID-Underflow-Report-{VERSION_TAG}.pdf",
        "Example Investigation Report",
        "USB APID 0x06 Undersized Telecommand",
    ),
    (
        SOURCE_ROOT / "reports" / "04-spacecanbus-spoofing-replay-report.md",
        REPORTS_DIR / f"FARI-Case-04-SpaceCAN-Spoofing-Replay-Report-{VERSION_TAG}.pdf",
        "Example Investigation Report",
        "SpaceCAN Spoofing and Replay",
    ),
    (
        SOURCE_ROOT / "reports" / "FARI-CONSOLIDATED-ASSESSMENT-REPORT.md",
        REPORTS_DIR / f"FARI-Consolidated-Assessment-Report-{VERSION_TAG}.pdf",
        "Example Consolidated Report",
        "Pilot Validation Portfolio",
    ),
    (
        SOURCE_ROOT / "reports" / "FARI-PROCESS-VALIDATION.md",
        REPORTS_DIR / f"FARI-Process-Validation-{VERSION_TAG}.pdf",
        "Framework Validation Report",
        "Process Validation and Observed Framework Behavior",
    ),
    (
        SOURCE_ROOT / "reports" / "FARI-SCENARIO-TEST-CATALOG-RELEASE.md",
        REPORTS_DIR / f"FARI-SCENARIO-TEST_CATALOG-{VERSION_TAG}.pdf",
        "Scenario Test Catalog",
        "Versioned Demo Portfolio and Coverage Matrix",
    ),
]


def copy_sources() -> None:
    if RELEASE_ROOT.exists():
        shutil.rmtree(RELEASE_ROOT)
    for relative, candidates in CURRENT_MARKDOWN_SOURCES.items():
        source = next((path for path in candidates if path.exists()), None)
        if source is None:
            raise FileNotFoundError(f"Missing source for release document: {relative}")
        destination = SOURCE_ROOT / relative
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, destination)
    legacy_dir = SOURCE_ROOT / "_archive" / "legacy-docx"
    legacy_dir.mkdir(parents=True, exist_ok=True)
    for source in LEGACY_DOCX_SOURCES:
        if source.exists():
            shutil.copy2(source, legacy_dir / source.name)
    combined = SOURCE_ROOT / "specification" / "FARI-SPECIFICATION-COMPLETE.md"
    spec = (SOURCE_ROOT / "specification" / "FARI-SPECIFICATION.md").read_text(encoding="utf-8").strip()
    combined.write_text(
        spec
        + "\n",
        encoding="utf-8",
    )
    _build_catalog_release_markdown()


def _build_catalog_release_markdown() -> None:
    source = SOURCE_ROOT / "reports" / "FARI-SCENARIO-TEST-CATALOG.md"
    text = source.read_text(encoding="utf-8")
    lines = text.splitlines()

    def section(title: str) -> str:
        pattern = rf"^## {re.escape(title)}\n(?P<body>.*?)(?=^## |\Z)"
        match = re.search(pattern, text, re.MULTILINE | re.DOTALL)
        return match.group("body").strip() if match else ""

    coverage = section("Coverage summary")
    filters = section("Filter expectations")
    demo = section("Suggested demo order")

    table_rows: list[list[str]] = []
    capture = False
    for line in lines:
        if line.startswith("| Assessment |"):
            capture = True
            continue
        if capture:
            if not line.startswith("|"):
                break
            if line.startswith("| ---"):
                continue
            table_rows.append([cell.strip() for cell in line.strip().strip("|").split("|")])

    release_lines = [
        "# FARI Scenario Test Catalog",
        "",
        "Seeded workspace catalog for the local FARI demo database.",
        "",
        "This release edition keeps the same portfolio coverage as the canonical source,",
        "but restructures the assessment list for easier executive reading in PDF form.",
        "",
        "## Coverage summary",
        "",
        coverage,
        "",
        "## Assessment portfolio",
        "",
    ]
    for row in table_rows:
        if len(row) < 5:
            continue
        assessment, result, disposition, versions, notes = row[:5]
        release_lines.extend(
            [
                f"### {assessment}",
                "",
                f"- Current result: `{result}`",
                f"- Scenario disposition: `{disposition}`",
                f"- Versions: {versions.replace('<br>', ' / ')}",
                f"- Notes: {notes}",
                "",
            ]
        )
    release_lines.extend(
        [
            "## Filter expectations",
            "",
            filters,
            "",
            "## Suggested demo order",
            "",
            demo,
            "",
        ]
    )
    (SOURCE_ROOT / "reports" / "FARI-SCENARIO-TEST-CATALOG-RELEASE.md").write_text(
        "\n".join(release_lines),
        encoding="utf-8",
    )


def _cover(doc: Document, label: str, title: str, subtitle: str) -> None:
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("FARI")
    set_run_font(run, size=30, color=INK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(label)
    set_run_font(run, size=18, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(title)
    set_run_font(run, size=15, color=INK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(subtitle)
    set_run_font(run, size=11, color=MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    run = p.add_run(f"{FARI_GENERATED_WITH}  |  {date.today().isoformat()}")
    set_run_font(run, size=10, color=MUTED, bold=True)

    for _ in range(3):
        doc.add_paragraph()

    add_callout(
        doc,
        "This is the versioned release edition of the FARI documentation bundle. "
        "Canonical working sources for this release are stored under the same version folder.",
    )
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _report_cover(doc: Document, label: str, title: str, source_name: str, conclusion: str | None) -> None:
    _cover(doc, label, title, f"Release source: {source_name}")
    if conclusion:
        add_result_banner(doc, conclusion, "Qualitative executive conclusion from the example report")


def _make_doc() -> Document:
    doc = Document()
    configure_styles(doc)
    configure_running_furniture(doc.sections[0])
    return doc


def _extract_conclusion(text: str) -> str | None:
    patterns = [
        r"\|\s*(?:Overall )?Conclusion\s*\|\s*\*\*([^*]+)\*\*\s*\|",
        r"\*\*Conclusion:\*\*\s*([^\n]+)",
    ]
    for pattern in patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            return match.group(1).strip()
    return None


def build_specification_docx() -> Path:
    BUILD_ROOT.mkdir(parents=True, exist_ok=True)
    doc = _make_doc()
    _cover(
        doc,
        "Framework Specification",
        f"Framework for Aerospace Research and Investigation {VERSION_TAG}",
        "Canonical specification, glossary, workflow, and release guidance",
    )
    source = SOURCE_ROOT / "specification" / "FARI-SPECIFICATION-COMPLETE.md"
    markdown_to_doc(doc, source, page_break=False)
    output = BUILD_ROOT / f"FARI-Specification-{VERSION_TAG}.docx"
    doc.save(output)
    return output


def build_manual_template_docx() -> Path:
    doc = _make_doc()
    _cover(
        doc,
        "Manual Fill Template",
        f"FARI Manual Template {VERSION_TAG}",
        "Report-author worksheet aligned with the finalized workflow and field names",
    )
    source = SOURCE_ROOT / "templates" / "FARI-Manual-Template.md"
    markdown_to_doc(doc, source, page_break=False)
    output = BUILD_ROOT / f"FARI-Manual-Template-{VERSION_TAG}.docx"
    doc.save(output)
    return output


def build_report_docx(source: Path, label: str, title: str) -> Path:
    doc = _make_doc()
    source_text = source.read_text(encoding="utf-8")
    conclusion = _extract_conclusion(source_text)
    _report_cover(doc, label, title, source.name, conclusion)
    markdown_to_doc(doc, source, page_break=False)
    output = BUILD_ROOT / (source.stem + ".docx")
    doc.save(output)
    return output


def main() -> None:
    copy_sources()
    SPEC_DIR.mkdir(parents=True, exist_ok=True)
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    QA_ROOT.mkdir(parents=True, exist_ok=True)
    print(build_specification_docx())
    print(build_manual_template_docx())
    for source, _pdf_target, label, title in REPORT_BUILD_PLAN:
        print(build_report_docx(source, label, title))


if __name__ == "__main__":
    main()
