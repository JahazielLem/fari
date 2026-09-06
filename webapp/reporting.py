"""Word report generation for web-created FARI assessments."""

from __future__ import annotations

from collections import Counter
from contextvars import ContextVar
from datetime import date
from html import escape
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from tools.build_reference_docx import (
    BLUE,
    INK,
    MUTED,
    add_callout,
    configure_running_furniture,
    configure_styles,
    set_cell_shading,
    set_repeat_table_header,
    set_run_font,
    set_table_geometry,
)
from webapp.i18n import load_catalog
from webapp.meta import FARI_GENERATED_WITH


RESULT_STYLE = {
    "meets": ("MEETS", "84ABF8", "FFFFFF"),
    "does_not_meet": ("DOES NOT MEET", "B42318", "FFFFFF"),
    "inconclusive": ("INCONCLUSIVE", "B54708", "FFFFFF"),
    "not_assessed": ("NOT ASSESSED", "475467", "FFFFFF"),
}

_REPORT_LANGUAGE = ContextVar("fari_report_language", default="en")
_REPORT_CATALOG = load_catalog(str(Path(__file__).resolve().parent / "translations"))


def report_text(key: str, **values) -> str:
    return _REPORT_CATALOG.get(f"report.{key}", _REPORT_LANGUAGE.get(), **values)


def human(value: str) -> str:
    raw = "" if value is None else str(value)
    return _REPORT_CATALOG.human(raw, _REPORT_LANGUAGE.get())


def revision_label(assessment) -> str:
    try:
        count = assessment["revision_count"]
    except (KeyError, TypeError):
        count = getattr(assessment, "revision_count", 0)
    return f"v{count}" if count else "v0"


def add_cover(
    doc,
    label: str,
    title: str,
    client_name: str,
    report_author: str,
    document_state: str,
    is_draft: bool = False,
):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("FARI")
    set_run_font(run, size=31, color=INK, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(label)
    set_run_font(run, size=18, color=BLUE, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run(title)
    set_run_font(run, size=15, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        f"{report_text('client')}: {client_name}  |  "
        f"{report_text('report_author')}: {report_author}"
    )
    set_run_font(run, size=10.5, color=MUTED, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(date.today().isoformat())
    set_run_font(run, size=10, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{report_text('generated_with')} {FARI_GENERATED_WITH}")
    set_run_font(run, size=9.5, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(document_state)
    set_run_font(
        run,
        size=16,
        color="B42318" if is_draft else "15803D",
        bold=True,
    )
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_heading(doc, text: str, level: int = 1):
    p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
    p.add_run(text)
    return p


def add_result_banner(doc, conclusion: str, confidence: str, action: str):
    _label, fill, text_color = RESULT_STYLE.get(
        conclusion, RESULT_STYLE["not_assessed"]
    )
    label = report_text(f"result_{conclusion}")
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(label)
    set_run_font(run, size=20, color=text_color, bold=True)
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(
        f"{report_text('confidence')}: {human(confidence)}  |  "
        f"{report_text('required_action')}: {human(action)}"
    )
    set_run_font(run, size=10.5, color=text_color, bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_key_value_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [2700, 6660])
    set_repeat_table_header(table.rows[0])
    table.cell(0, 0).text = report_text("field")
    table.cell(0, 1).text = report_text("result")
    for cell in table.rows[0].cells:
        set_cell_shading(cell, "E8EEF5")
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=9.5, color=INK, bold=True)
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value or report_text("not_supplied"))
        for run in cells[0].paragraphs[0].runs:
            set_run_font(run, size=9.2, color=INK, bold=True)
        for run in cells[1].paragraphs[0].runs:
            set_run_font(run, size=9.2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_list_section(doc, title: str, text: str, level: int = 2):
    add_heading(doc, title, level)
    items = [item.strip() for item in (text or "").splitlines() if item.strip()]
    if not items:
        p = doc.add_paragraph()
        run = p.add_run(f"{report_text('not_supplied')}.")
        set_run_font(run, size=10.5, color=MUTED, italic=True)
        return
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_findings(doc, findings):
    add_heading(doc, report_text("normalized_findings"), 2)
    if not findings:
        add_callout(doc, report_text("no_normalized_findings"))
        return
    for finding in findings:
        add_heading(doc, f"{finding['fari_id']}: {finding['title']}", 3)
        add_key_value_table(
            doc,
            [
                (report_text("state"), human(finding["state"])),
                (report_text("condition"), finding["condition_text"]),
                (report_text("observed_effect"), finding["observed_effect"]),
                (report_text("credible_impact"), finding["credible_impact"]),
                (report_text("mapping_state"), human(finding["mapping_state"])),
                (report_text("mapping_rationale"), finding["mapping_rationale"]),
            ],
        )


def build_investigation_report(assessment, investigation, asset, evidence, findings, language="en"):
    _REPORT_LANGUAGE.set(language)
    document_state = "FINAL" if investigation["status"] == "closed" else "DRAFT"
    doc = Document()
    configure_styles(doc)
    configure_running_furniture(doc.sections[0])
    add_cover(
        doc,
        report_text("investigation_report"),
        investigation["title"],
        assessment["client_name"],
        assessment["report_author"],
        report_text("final") if document_state == "FINAL" else report_text("draft"),
        is_draft=document_state == "DRAFT",
    )
    if document_state == "DRAFT":
        add_callout(
            doc,
            report_text("draft_investigation"),
        )
    add_result_banner(
        doc,
        investigation["conclusion"],
        investigation["confidence"],
        investigation["required_action"],
    )
    add_heading(doc, f"1. {report_text('fari_result_card')}", 1)
    add_key_value_table(
        doc,
        [
            (report_text("overall_conclusion"), human(investigation["conclusion"])),
            (report_text("scenario_disposition"), human(investigation["scenario_state"])),
            (report_text("confidence"), human(investigation["confidence"])),
            (report_text("required_action"), human(investigation["required_action"])),
            (report_text("generated_with_field"), FARI_GENERATED_WITH),
            (report_text("assessment_revision"), revision_label(assessment)),
            (report_text("priority"), human(investigation["priority"])),
            (report_text("scope_boundary"), investigation["scope_boundary"]),
            (report_text("rationale"), investigation["rationale"]),
        ],
    )
    add_heading(doc, f"2. {report_text('frame')}", 1)
    add_key_value_table(
        doc,
        [
            (report_text("assessment"), assessment["fari_id"]),
            (report_text("assessment_revision"), revision_label(assessment)),
            (report_text("client"), assessment["client_name"]),
            (report_text("investigation"), investigation["fari_id"]),
            (report_text("claim"), investigation["claim_id"]),
            (report_text("claim_description"), investigation["claim_description"]),
            (report_text("gating_claim"), report_text("yes") if investigation["gating"] else report_text("no")),
            (report_text("asset"), asset["fari_id"] + " - " + asset["name"] if asset else report_text("not_supplied")),
            (report_text("technical_reporter"), investigation["technical_reporter"]),
            (report_text("report_author"), assessment["report_author"]),
            (report_text("method"), investigation["method"]),
            (report_text("environment"), investigation["environment"]),
        ],
    )
    add_heading(doc, f"3. {report_text('acquire')}", 1)
    add_list_section(doc, report_text("evidence_backed_facts"), investigation["facts"])
    add_list_section(doc, report_text("evidence_producer_assertions"), investigation["assertions"])
    add_list_section(doc, report_text("report_author_inferences"), investigation["inferences"])
    add_list_section(doc, report_text("assumptions"), investigation["assumptions"])
    add_list_section(doc, report_text("contradictions"), investigation["contradictions"])
    add_list_section(doc, report_text("missing_information"), investigation["gaps"])
    add_heading(doc, report_text("evidence_sufficiency"), 2)
    add_key_value_table(
        doc,
        [
            (report_text("technical_condition"), human(investigation["technical_sufficiency"])),
            (report_text("reachability"), human(investigation["reachability_sufficiency"])),
            (report_text("mission_consequence"), human(investigation["mission_sufficiency"])),
        ],
    )
    add_heading(doc, report_text("evidence_index"), 2)
    if evidence:
        add_key_value_table(
            doc,
            [
                (
                    item["fari_id"],
                    f"<strong>{item['description']}</strong> | {item['filename']} <br>SHA-256 {item['sha256']}",
                )
                for item in evidence
            ],
        )
    else:
        add_callout(doc, report_text("no_evidence_files"))
    add_heading(doc, f"4. {report_text('relate')}", 1)
    add_findings(doc, findings)
    add_heading(doc, f"5. {report_text('inform')}", 1)
    add_list_section(doc, report_text("recommendations"), investigation["recommendations"])
    add_list_section(doc, report_text("acceptance_criteria"), investigation["acceptance_criteria"])
    add_heading(doc, report_text("closing_result"), 1)
    add_result_banner(
        doc,
        investigation["conclusion"],
        investigation["confidence"],
        investigation["required_action"],
    )
    doc.core_properties.author = assessment["report_author"]
    doc.core_properties.title = f"{investigation['fari_id']} {investigation['title']}"
    doc.core_properties.subject = FARI_GENERATED_WITH
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output


def overall_conclusion(investigations) -> str:
    gating = [row for row in investigations if row["gating"]]
    if not gating:
        return "not_assessed"
    values = {row["conclusion"] for row in gating}
    if "does_not_meet" in values:
        return "does_not_meet"
    if "inconclusive" in values or "not_assessed" in values:
        return "inconclusive"
    return "meets"


def overall_confidence(investigations) -> str:
    values = {row["confidence"] for row in investigations}
    if "low" in values:
        return "low"
    if "medium" in values:
        return "medium"
    return "high" if values else "low"


def consolidated_action(conclusion: str) -> str:
    return {
        "meets": "accept",
        "does_not_meet": "remediate",
        "inconclusive": "extend_investigation",
        "not_assessed": "extend_investigation",
    }[conclusion]


def build_consolidated_report(assessment, assets, investigations, language="en"):
    _REPORT_LANGUAGE.set(language)
    document_state = "FINAL" if assessment["status"] == "closed" else "DRAFT"
    conclusion = overall_conclusion(investigations)
    confidence = overall_confidence(investigations)
    action = consolidated_action(conclusion)
    doc = Document()
    configure_styles(doc)
    configure_running_furniture(doc.sections[0])
    add_cover(
        doc,
        report_text("consolidated_report"),
        assessment["title"],
        assessment["client_name"],
        assessment["report_author"],
        report_text("final") if document_state == "FINAL" else report_text("draft"),
        is_draft=document_state == "DRAFT",
    )
    if document_state == "DRAFT":
        add_callout(
            doc,
            report_text("draft_assessment"),
        )
    add_result_banner(doc, conclusion, confidence, action)
    add_heading(doc, f"1. {report_text('overall_fari_result_card')}", 1)
    add_key_value_table(
        doc,
        [
            (report_text("overall_conclusion"), human(conclusion)),
            (report_text("confidence"), human(confidence)),
            (report_text("required_action"), human(action)),
            (report_text("generated_with_field"), FARI_GENERATED_WITH),
            (report_text("assessment_revision"), revision_label(assessment)),
            (report_text("assessment"), assessment["fari_id"]),
            (report_text("client"), assessment["client_name"]),
            (report_text("scope_boundary"), assessment["scope"]),
            (
                report_text("rationale"),
                report_text("derived_from_gating"),
            ),
        ],
    )
    add_heading(doc, f"2. {report_text('frame_client_context')}", 1)
    add_key_value_table(
        doc,
        [
            (report_text("mission_context"), assessment["mission_context"]),
            (report_text("scope"), assessment["scope"]),
            (report_text("authorization"), assessment["authorization"]),
            (report_text("exclusions"), assessment["exclusions"]),
            (report_text("report_author"), assessment["report_author"]),
        ],
    )
    add_heading(doc, f"3. {report_text('asset_register')}", 1)
    add_key_value_table(
        doc,
        [
            (
                asset["fari_id"],
                f"{asset['name']} | {human(asset['segment'])} | "
                f"{human(asset['coverage'])} | {asset['description']}",
            )
            for asset in assets
        ]
        or [(report_text("asset"), report_text("no_assets"))],
    )
    add_heading(doc, f"4. {report_text('investigation_results')}", 1)
    for investigation in investigations:
        add_heading(
            doc,
            f"{investigation['fari_id']}: {investigation['title']}",
            2,
        )
        add_result_banner(
            doc,
            investigation["conclusion"],
            investigation["confidence"],
            investigation["required_action"],
        )
        add_key_value_table(
            doc,
            [
                (report_text("claim"), f"{investigation['claim_id']}: {investigation['claim_description']}"),
                (report_text("gating"), report_text("yes") if investigation["gating"] else report_text("no")),
                (report_text("priority"), human(investigation["priority"])),
                (report_text("scenario_disposition"), human(investigation["scenario_state"])),
                (report_text("technical_sufficiency"), human(investigation["technical_sufficiency"])),
                (report_text("reachability_sufficiency"), human(investigation["reachability_sufficiency"])),
                (report_text("mission_sufficiency"), human(investigation["mission_sufficiency"])),
                (report_text("rationale"), investigation["rationale"]),
            ],
        )
        add_heading(doc, report_text("evidence_index"), 3)
        if investigation.evidence:
            add_key_value_table(
                doc,
                [
                    (
                        item["fari_id"],
                        f"{item['filename']} | SHA-256 {item['sha256']} | {item['description']}",
                    )
                    for item in investigation.evidence
                ],
            )
        else:
            add_callout(doc, report_text("no_evidence_files"))
        add_findings(doc, investigation.findings)
    add_heading(doc, f"5. {report_text('totalized_results')}", 1)
    counts = Counter(row["conclusion"] for row in investigations)
    add_key_value_table(
        doc,
        [
            (RESULT_STYLE[key][0], counts.get(key, 0))
            for key in ("meets", "does_not_meet", "inconclusive", "not_assessed")
        ],
    )
    scenario_counts = Counter(row["scenario_state"] for row in investigations)
    add_heading(doc, report_text("scenario_dispositions"), 2)
    add_key_value_table(
        doc,
        [
            (human(key), scenario_counts.get(key, 0))
            for key in ("demonstrated", "plausible", "not_demonstrated", "not_evaluated")
        ],
    )
    add_heading(doc, f"6. {report_text('required_actions')}", 1)
    for investigation in investigations:
        add_heading(doc, investigation["fari_id"], 2)
        add_list_section(doc, report_text("recommendations"), investigation["recommendations"], 3)
        add_list_section(doc, report_text("acceptance_criteria"), investigation["acceptance_criteria"], 3)
    add_heading(doc, report_text("closing_overall_result"), 1)
    add_result_banner(doc, conclusion, confidence, action)
    doc.core_properties.author = assessment["report_author"]
    doc.core_properties.title = f"{assessment['fari_id']} Consolidated Assessment"
    doc.core_properties.subject = FARI_GENERATED_WITH
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output


def build_investigation_pdf(assessment, investigation, asset, evidence, findings, language="en"):
    """Build the immutable final investigation report as a PDF."""
    _REPORT_LANGUAGE.set(language)
    output = BytesIO()
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="FariTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=24,
            leading=28,
            textColor=colors.HexColor("#4C4F69"),
            alignment=TA_CENTER,
            spaceAfter=14,
        )
    )
    styles.add(
        ParagraphStyle(
            name="FariHeading",
            parent=styles["Heading2"],
            textColor=colors.HexColor("#1E66F5"),
            spaceBefore=14,
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            name="FariSmall",
            parent=styles["BodyText"],
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#6C6F85"),
        )
    )

    doc = SimpleDocTemplate(
        output,
        pagesize=letter,
        leftMargin=0.65 * inch,
        rightMargin=0.65 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.7 * inch,
        title=f"{investigation['fari_id']} {investigation['title']}",
        author=assessment["report_author"],
    )
    story = [
        Spacer(1, 1.1 * inch),
        Paragraph(
            f"FARI<br/>{escape(report_text('final_investigation'))}<br/>{escape(report_text('report_label'))}",
            styles["FariHeading"],
        ),
        Paragraph(escape(investigation["title"]), styles["FariTitle"]),
        Paragraph(
            f"{escape(report_text('client'))}: {escape(assessment['client_name'])}<br/>"
            f"{escape(report_text('report_author'))}: {escape(assessment['report_author'])}<br/>"
            f"{escape(report_text('closed_report_generated'))}: {date.today().isoformat()}<br/>"
            f"{escape(report_text('generated_with'))} {escape(FARI_GENERATED_WITH)}",
            styles["BodyText"],
        ),
        Spacer(1, 0.3 * inch),
        _pdf_result_banner(investigation, styles),
        PageBreak(),
        Paragraph(report_text("fari_result_card"), styles["FariHeading"]),
        _pdf_key_value(
            [
                (report_text("overall_conclusion"), human(investigation["conclusion"])),
                (report_text("scenario_disposition"), human(investigation["scenario_state"])),
                (report_text("confidence"), human(investigation["confidence"])),
                (report_text("required_action"), human(investigation["required_action"])),
                (report_text("generated_with_field"), FARI_GENERATED_WITH),
                (report_text("assessment_revision"), revision_label(assessment)),
                (report_text("priority"), human(investigation["priority"])),
                (report_text("scope_boundary"), investigation["scope_boundary"]),
                (report_text("rationale"), investigation["rationale"]),
            ],
            styles,
        ),
        Paragraph(report_text("frame"), styles["FariHeading"]),
        _pdf_key_value(
            [
                (report_text("assessment"), assessment["fari_id"]),
                (report_text("assessment_revision"), revision_label(assessment)),
                (report_text("investigation"), investigation["fari_id"]),
                (report_text("claim"), investigation["claim_id"]),
                (report_text("claim_description"), investigation["claim_description"]),
                (report_text("asset"), f"{asset['fari_id']} - {asset['name']}" if asset else report_text("not_supplied")),
                (report_text("technical_reporter"), investigation["technical_reporter"]),
                (report_text("method"), investigation["method"]),
                (report_text("environment"), investigation["environment"]),
            ],
            styles,
        ),
        Paragraph(report_text("acquire"), styles["FariHeading"]),
    ]
    for title, text in [
        (report_text("evidence_backed_facts"), investigation["facts"]),
        (report_text("evidence_producer_assertions"), investigation["assertions"]),
        (report_text("report_author_inferences"), investigation["inferences"]),
        (report_text("assumptions"), investigation["assumptions"]),
        (report_text("contradictions"), investigation["contradictions"]),
        (report_text("missing_information"), investigation["gaps"]),
    ]:
        story.extend(_pdf_list(title, text, styles))
    story.extend(
        [
            Paragraph(report_text("evidence_index"), styles["FariHeading"]),
            _pdf_key_value(
                [
                    (
                        item["fari_id"],
                        f"{item['filename']} | SHA-256 {item['sha256']} | {item['description']}",
                    )
                    for item in evidence
                ]
                or [(report_text("evidence"), report_text("no_evidence_uploaded"))],
                styles,
            ),
            Paragraph(f"{report_text('relate')}: {report_text('normalized_findings')}", styles["FariHeading"]),
        ]
    )
    for finding in findings:
        story.append(Paragraph(f"{escape(finding['fari_id'])}: {escape(finding['title'])}", styles["Heading3"]))
        story.append(
            _pdf_key_value(
                [
                    (report_text("state"), human(finding["state"])),
                    (report_text("condition"), finding["condition_text"]),
                    (report_text("observed_effect"), finding["observed_effect"]),
                    (report_text("credible_impact"), finding["credible_impact"]),
                    (report_text("mapping_state"), human(finding["mapping_state"])),
                    (report_text("mapping_rationale"), finding["mapping_rationale"]),
                ],
                styles,
            )
        )
    story.extend(
        [
            Paragraph(report_text("inform"), styles["FariHeading"]),
            *_pdf_list(report_text("recommendations"), investigation["recommendations"], styles),
            *_pdf_list(report_text("acceptance_criteria"), investigation["acceptance_criteria"], styles),
            Spacer(1, 0.2 * inch),
            _pdf_result_banner(investigation, styles),
        ]
    )
    doc.build(story, onFirstPage=_pdf_footer, onLaterPages=_pdf_footer)
    output.seek(0)
    return output


def build_consolidated_pdf(assessment, assets, investigations, language="en"):
    """Build the immutable final assessment-level report as a PDF."""
    _REPORT_LANGUAGE.set(language)
    output = BytesIO()
    styles = _pdf_styles()
    conclusion = overall_conclusion(investigations)
    confidence = overall_confidence(investigations)
    action = consolidated_action(conclusion)
    doc = SimpleDocTemplate(
        output,
        pagesize=letter,
        leftMargin=0.65 * inch,
        rightMargin=0.65 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.7 * inch,
        title=f"{assessment['fari_id']} Consolidated Assessment",
        author=assessment["report_author"],
    )
    overall = {
        "conclusion": conclusion,
        "confidence": confidence,
        "required_action": action,
    }
    story = [
        Paragraph(
            f"FARI<br/>{escape(report_text('final_consolidated'))}<br/>{escape(report_text('assessment_label'))}",
            styles["FariHeading"],
        ),
        Spacer(1, 1.1 * inch),
        Paragraph(escape(assessment["title"]), styles["FariTitle"]),
        Paragraph(
            f"{escape(report_text('client'))}: {escape(assessment['client_name'])}<br/>"
            f"{escape(report_text('report_author'))}: {escape(assessment['report_author'])}<br/>"
            f"{escape(report_text('final_report_generated'))}: {date.today().isoformat()}<br/>"
            f"{escape(report_text('generated_with'))} {escape(FARI_GENERATED_WITH)}",
            styles["BodyText"],
        ),
        Spacer(1, 0.3 * inch),
        _pdf_result_banner(overall, styles),
        PageBreak(),
        Paragraph(report_text("overall_fari_result_card"), styles["FariHeading"]),
        _pdf_key_value(
            [
                (report_text("overall_conclusion"), human(conclusion)),
                (report_text("confidence"), human(confidence)),
                (report_text("required_action"), human(action)),
                (report_text("generated_with_field"), FARI_GENERATED_WITH),
                (report_text("assessment_revision"), revision_label(assessment)),
                (report_text("assessment"), assessment["fari_id"]),
                (report_text("scope_boundary"), assessment["scope"]),
                (report_text("rationale"), report_text("derived_from_gating")),
            ],
            styles,
        ),
        Paragraph(report_text("frame_client_context"), styles["FariHeading"]),
        _pdf_key_value(
            [
                (report_text("mission_context"), assessment["mission_context"]),
                (report_text("scope"), assessment["scope"]),
                (report_text("authorization"), assessment["authorization"]),
                (report_text("exclusions"), assessment["exclusions"]),
            ],
            styles,
        ),
        Paragraph(report_text("asset_register"), styles["FariHeading"]),
        _pdf_key_value(
            [
                (
                    asset["fari_id"],
                    f"{asset['name']} | {human(asset['segment'])} | "
                    f"{human(asset['coverage'])} | {asset['description']}",
                )
                for asset in assets
            ]
            or [(report_text("asset"), report_text("no_assets"))],
            styles,
        ),
        PageBreak(),
        Paragraph(report_text("investigation_results"), styles["FariHeading"]),
    ]
    for investigation in investigations:
        story.extend(
            [
                Paragraph(
                    f"{escape(investigation['fari_id'])}: {escape(investigation['title'])}",
                    styles["Heading2"],
                ),
                _pdf_result_banner(investigation, styles),
                _pdf_key_value(
                    [
                        (report_text("claim"), f"{investigation['claim_id']}: {investigation['claim_description']}"),
                        (report_text("gating"), report_text("yes") if investigation["gating"] else report_text("no")),
                        (report_text("scenario_disposition"), human(investigation["scenario_state"])),
                        (report_text("technical_sufficiency"), human(investigation["technical_sufficiency"])),
                        (report_text("reachability_sufficiency"), human(investigation["reachability_sufficiency"])),
                        (report_text("mission_sufficiency"), human(investigation["mission_sufficiency"])),
                        (report_text("scope_boundary"), investigation["scope_boundary"]),
                        (report_text("rationale"), investigation["rationale"]),
                    ],
                    styles,
                ),
                Paragraph(report_text("evidence_index"), styles["Heading3"]),
                _pdf_key_value(
                    [
                        (
                            item["fari_id"],
                            f"{item['filename']} | SHA-256 {item['sha256']} | {item['description']}",
                        )
                        for item in investigation.evidence
                    ]
                    or [(report_text("evidence"), report_text("no_evidence_uploaded"))],
                    styles,
                ),
                Paragraph(report_text("normalized_findings"), styles["Heading3"]),
            ]
        )
        for finding in investigation.findings:
            story.append(
                _pdf_key_value(
                    [
                        (report_text("finding_title"), f"{finding['fari_id']}: {finding['title']}"),
                        (report_text("state"), human(finding["state"])),
                        (report_text("condition"), finding["condition_text"]),
                        (report_text("observed_effect"), finding["observed_effect"]),
                        (report_text("credible_impact"), finding["credible_impact"]),
                        (report_text("mapping_state"), human(finding["mapping_state"])),
                        (report_text("mapping_rationale"), finding["mapping_rationale"]),
                    ],
                    styles,
                )
            )
        if not investigation.findings:
            story.append(Paragraph(report_text("no_findings"), styles["FariSmall"]))
        story.extend(
            [
                *_pdf_list(report_text("recommendations"), investigation["recommendations"], styles),
                *_pdf_list(report_text("acceptance_criteria"), investigation["acceptance_criteria"], styles),
                Spacer(1, 0.15 * inch),
            ]
        )

    conclusion_counts = Counter(row["conclusion"] for row in investigations)
    scenario_counts = Counter(row["scenario_state"] for row in investigations)
    story.extend(
        [
            PageBreak(),
            Paragraph(report_text("totalized_results"), styles["FariHeading"]),
            _pdf_key_value(
                [
                    (RESULT_STYLE[key][0], conclusion_counts.get(key, 0))
                    for key in ("meets", "does_not_meet", "inconclusive", "not_assessed")
                ],
                styles,
            ),
            Paragraph(report_text("scenario_dispositions"), styles["FariHeading"]),
            _pdf_key_value(
                [
                    (human(key), scenario_counts.get(key, 0))
                    for key in ("demonstrated", "plausible", "not_demonstrated", "not_evaluated")
                ],
                styles,
            ),
            Paragraph(report_text("closing_overall_result"), styles["FariHeading"]),
            _pdf_result_banner(overall, styles),
        ]
    )
    doc.build(story, onFirstPage=_pdf_footer, onLaterPages=_pdf_footer)
    output.seek(0)
    return output


def _pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="FariTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=24,
            leading=28,
            textColor=colors.HexColor("#4C4F69"),
            alignment=TA_CENTER,
            spaceAfter=14,
        )
    )
    styles.add(
        ParagraphStyle(
            name="FariHeading",
            parent=styles["Heading2"],
            textColor=colors.HexColor("#1E66F5"),
            alignment=TA_CENTER,
            spaceBefore=14,
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            name="FariSmall",
            parent=styles["BodyText"],
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#6C6F85"),
        )
    )
    return styles


def _pdf_result_banner(investigation, styles):
    _label, fill, _text_color = RESULT_STYLE.get(
        investigation["conclusion"], RESULT_STYLE["not_assessed"]
    )
    label = report_text(f"result_{investigation['conclusion']}")
    table = Table(
        [
            [Paragraph(f"<b>{label}</b>", styles["Title"])],
            [
                Paragraph(
                    f"{report_text('confidence')}: {human(investigation['confidence'])} | "
                    f"{report_text('required_action')}: {human(investigation['required_action'])}",
                    styles["BodyText"],
                )
            ],
        ],
        colWidths=[7.1 * inch],
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{fill}")),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 10),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
            ]
        )
    )
    return table


def _pdf_key_value(rows, styles):
    data = [
        [
            Paragraph(f"<b>{escape(report_text('field'))}</b>", styles["BodyText"]),
            Paragraph(f"<b>{escape(report_text('result'))}</b>", styles["BodyText"]),
        ]
    ]
    for label, value in rows:
        data.append(
            [
                Paragraph(f"<b>{escape(str(label))}</b>", styles["BodyText"]),
                Paragraph(escape(str(value or report_text("not_supplied"))).replace("\n", "<br/>"), styles["BodyText"]),
            ]
        )
    table = Table(data, colWidths=[1.6 * inch, 5.5 * inch], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6E9EF")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#BCC0CC")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return table


def _pdf_list(title, text, styles):
    items = [item.strip() for item in (text or "").splitlines() if item.strip()]
    result = [Paragraph(escape(title), styles["Heading3"])]
    if not items:
        result.append(Paragraph(f"{escape(report_text('not_supplied'))}.", styles["FariSmall"]))
    else:
        result.extend(
            Paragraph(f"• {escape(item)}", styles["BodyText"]) for item in items
        )
    return result


def _pdf_footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#6C6F85"))
    canvas.drawString(
        0.65 * inch,
        0.38 * inch,
        f"{report_text('framework_footer')} | {FARI_GENERATED_WITH}",
    )
    canvas.drawRightString(7.85 * inch, 0.38 * inch, f"{report_text('page')} {doc.page}")
    canvas.restoreState()
